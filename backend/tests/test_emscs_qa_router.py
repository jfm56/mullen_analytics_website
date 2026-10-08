"""EMSCS QA router integration test (flag gating + create->review->approve->export).

Uses a standalone app with the real router and a managed session (the shared db
fixture rolls back and does not handle commits, which the router performs). Auth is
overridden to a reviewer ctx — the dual-mode auth wrappers reuse the already-tested
require_membership/session paths; this test targets the endpoint→service→export wiring."""
import io
import os
import uuid

import openpyxl
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, text
from sqlalchemy.orm import sessionmaker

from app.config import get_settings  # noqa: E402

from app.database import Base  # noqa: E402
from app.models import emscs_qa  # noqa: F401,E402 — register QA tables
from app.models.agency import Agency  # noqa: E402
from app.models.user import User  # noqa: E402
from app.services.auth import hash_password  # noqa: E402
from app.routers import emscs_qa as qa  # noqa: E402

_ENGINE = create_engine(os.environ.get("DATABASE_URL", "postgresql://postgres:postgres@localhost:5432/mullen_analytics"))
_Session = sessionmaker(bind=_ENGINE)

CHART = {
    "external_ref": "RTR-1", "transported": True, "age_years": 70, "primary_impression": "Chest pain",
    "vitals": [{"sbp": 120, "bp_method": "auto", "pain": 0}],
    "securement_text": "all straps", "signatures": {"patient": True, "crew": True},
    "structured_fields": {"airway": {"charted": "i-gel", "narrative": "ALS intubation"}},
    "add_actions": [{"kind": "medication", "name": "Aspirin"}],
}


@pytest.fixture(scope="module")
def qa_client():
    os.environ["EMSCS_QA_V1_ENABLED"] = "true"
    get_settings.cache_clear()
    Base.metadata.create_all(_ENGINE)
    s = _Session()
    agency = Agency(agency_name="rtr", slug=f"rtr-{uuid.uuid4().hex[:8]}")
    actor = User(email=f"rtr-{uuid.uuid4().hex[:8]}@t.test", password_hash=hash_password("x"))
    s.add(agency); s.add(actor); s.commit()
    aid, uid = agency.id, actor.id

    app = FastAPI()
    app.include_router(qa.router, prefix="/api")

    def _access(agency_id):
        return qa.QaCtx(s, agency_id, uid, True, True)
    app.dependency_overrides[qa.qa_access] = _access

    yield TestClient(app), str(aid), s

    for t in ("qa_audit_events", "qa_scores", "qa_findings", "qa_indicator_reviews",
              "qa_crew_feedback", "qa_review_sessions", "qa_charts"):
        s.execute(text(f"DELETE FROM {t} WHERE agency_id = :a"), {"a": aid})  # nosec B608 - fixed table name
    s.execute(text("DELETE FROM agencies WHERE id = :a"), {"a": aid})
    s.execute(text("DELETE FROM users WHERE id = :u"), {"u": uid})
    s.commit(); s.close()
    os.environ.pop("EMSCS_QA_V1_ENABLED", None)
    get_settings.cache_clear()


def test_health_reflects_flag(qa_client):
    client, _aid, _s = qa_client
    assert client.get("/api/v1/qa/health").json()["enabled"] is True


def test_library_served(qa_client):
    client, _aid, _s = qa_client
    lib = client.get("/api/v1/qa/library").json()
    assert len(lib["indicators"]) == 84 and len(lib["domains"]) == 8


def test_create_review_review_and_export(qa_client):
    client, aid, _s = qa_client
    r = client.post(f"/api/v1/agencies/{aid}/qa/reviews", json={"chart": CHART})
    assert r.status_code == 201
    sid = r.json()["session_id"]

    detail = client.get(f"/api/v1/agencies/{aid}/qa/sessions/{sid}").json()
    assert len(detail["indicator_reviews"]) >= 12
    assert detail["findings"], "expected findings"
    assert all(f["evidence"] for f in detail["findings"] if f["type"] != "manual")

    # override one indicator (reason enforced)
    ir = next(i for i in detail["indicator_reviews"] if i["number"] == 3)
    assert client.post(f"/api/v1/agencies/{aid}/qa/indicators/{ir['id']}/override",
                       json={"verdict": "fail", "reason": ""}).status_code == 400
    assert client.post(f"/api/v1/agencies/{aid}/qa/indicators/{ir['id']}/override",
                       json={"verdict": "pass", "reason": "manual BP present on strip"}).status_code == 200
    d2 = client.get(f"/api/v1/agencies/{aid}/qa/sessions/{sid}").json()
    ir2 = next(i for i in d2["indicator_reviews"] if i["number"] == 3)
    assert ir2["human"]["verdict"] == "pass" and ir2["automated"]["verdict"] == "fail"  # automated preserved

    # resolve remaining indicators + acknowledge criticals + domain scores -> approve
    for i in d2["indicator_reviews"]:
        if i["number"] == 3:
            continue
        if (i["automated"] or {}).get("verdict") == "human_review_required":
            client.post(f"/api/v1/agencies/{aid}/qa/indicators/{i['id']}/override",
                        json={"verdict": "na", "reason": "n/a for this chart"})
        else:
            client.post(f"/api/v1/agencies/{aid}/qa/indicators/{i['id']}/accept")
    for f in d2["findings"]:
        if f["final_severity"] == "Critical":
            client.post(f"/api/v1/agencies/{aid}/qa/findings/{f['id']}/acknowledge")
    client.post(f"/api/v1/agencies/{aid}/qa/sessions/{sid}/domain-scores", json={"scores": [3, 3, 4, 2, 2, 3, 3, 4]})
    ap = client.post(f"/api/v1/agencies/{aid}/qa/sessions/{sid}/approve")
    assert ap.status_code == 200 and ap.json()["approved_tier"] is not None

    # audit trail populated
    audit = client.get(f"/api/v1/agencies/{aid}/qa/audit").json()["events"]
    assert any(e["action"] == "session_approved" for e in audit)
    assert any(e["action"] == "indicator_overridden" for e in audit)

    # export returns a valid xlsx with the 12 sheets
    ex = client.get(f"/api/v1/agencies/{aid}/qa/export")
    assert ex.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(ex.content))
    assert "Chart Log" in wb.sheetnames and "Dashboard" in wb.sheetnames
